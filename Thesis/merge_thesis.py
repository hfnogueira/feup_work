"""
Merge all thesis chapters using docxcompose — properly handles images.
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docxcompose.composer import Composer
import copy, os

THESIS_DIR = '/sessions/charming-modest-galileo/mnt/Thesis/'
OUT_PATH   = '/sessions/charming-modest-galileo/mnt/Thesis/thesis_complete.docx'

CHAPTERS = [
    ('Abstract',          ['abstract_v2.docx']),
    ('Introduction',      ['chapter1_introduction_v2.docx']),
    ('Literature Review', ['chapter2_literature_v2.docx', 'chapter2_expand.docx', 'chapter2_cycle_figure.docx']),
    ('Methodology',       ['chapter3_methods_v2.docx', 'chapter3_prediction_formula.docx', 'chapter3_stationarity.docx', 'chapter3_data_pipeline.docx', 'chapter3_acf.docx', 'chapter3_context_figures.docx']),
    ('Rule Extraction',   ['chapter4_v4.docx', 'chapter4_overlap_freq.docx']),
    ('Prediction',        ['chapter5_prediction_v2.docx', 'chapter5_per_scenario.docx', 'chapter5_expand.docx']),
    ('Discussion',        ['chapter6_discussion.docx', 'chapter6_expand.docx']),
    ('Conclusions',       ['chapter7_v3.docx']),
    ('Figures',           ['figures.docx']),
    ('References',        ['references_v2.docx']),
    ('Appendix A',        ['appendix_a_rules.docx']),
    ('Appendix B',        ['appendix_b_eda.docx']),
    ('Appendix C',        ['appendix_c_features.docx']),
    ('Appendix D',        ['appendix_d_feature_map.docx']),
]

CH3_INJECT_AFTER = 'This generates 18 evaluation scenarios for RDD'
CH3_SENTENCE = (
    'Specifically, with a minimum training fraction of 80% applied to 89 RDD observations, '
    'the minimum training window covers 71 years (1934–2004), leaving 18 subsequent years '
    '(2005–2022) as sequential test years — one per scenario. For RVV, 80% of 80 observations '
    'yields a minimum training window of 64 years (1942–2005), leaving 16 test years (2006–2021).'
)

def inject_ch3_sentence(doc):
    for i, para in enumerate(doc.paragraphs):
        if CH3_INJECT_AFTER in para.text:
            new_p = OxmlElement('w:p')
            new_r = OxmlElement('w:r')
            new_rPr = OxmlElement('w:rPr')
            rFonts = OxmlElement('w:rFonts')
            rFonts.set(qn('w:ascii'), 'Times New Roman')
            rFonts.set(qn('w:hAnsi'), 'Times New Roman')
            new_rPr.append(rFonts)
            sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '24')
            new_rPr.append(sz)
            new_r.append(new_rPr)
            t = OxmlElement('w:t'); t.text = CH3_SENTENCE
            new_r.append(t)
            new_p.append(new_r)
            para._p.addnext(new_p)
            print(f'  ✓ Ch3 sentence injected')
            return True
    print('  ⚠ Ch3 inject target not found')
    return False

def add_title_page(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('PORTO UNIVERSITY'); r.font.size = Pt(14); r.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Faculty of Engineering of the University of Porto'); r.font.size = Pt(12)

    for _ in range(4): doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Sub-Group Discovery and Meaningful Distribution Rules\nin Time Series for Wine Production Forecasting')
    r.font.size = Pt(18); r.font.bold = True

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('MSc Thesis in Data Science and Engineering'); r.font.size = Pt(13); r.font.italic = True

    for _ in range(5): doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Hugo Nogueira'); r.font.size = Pt(13); r.font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('up202103617'); r.font.size = Pt(12)

    for _ in range(3): doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Supervised by:'); r.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Prof. João Pedro Mendes Moreira (FEUP)'); r.font.size = Pt(12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Prof. Mário Campos Cunha (FCUP)'); r.font.size = Pt(12)

    for _ in range(3): doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('June 2026'); r.font.size = Pt(12)

# ── Build ─────────────────────────────────────────────────────────────────────
print('Building thesis_complete.docx (v2 with images)...')

# Start with title page as master
master = Document()
style = master.styles['Normal']
style.font.name = 'Times New Roman'; style.font.size = Pt(12)
section = master.sections[0]
section.page_width = Cm(21); section.page_height = Cm(29.7)
section.left_margin = Cm(2.5); section.right_margin = Cm(2.0)
section.top_margin = Cm(2.5); section.bottom_margin = Cm(2.0)

print('Adding title page...')
add_title_page(master)

composer = Composer(master)

for label, files in CHAPTERS:
    print(f'\nChapter: {label}')
    for fname in files:
        fpath = os.path.join(THESIS_DIR, fname)
        if not os.path.exists(fpath):
            print(f'  MISSING: {fname}')
            continue
        print(f'  Appending {fname}...')

        src = Document(fpath)

        # Inject Ch3 sentence
        if fname == 'chapter3_methods_v2.docx':
            inject_ch3_sentence(src)

        # Add page break before first file of each chapter
        if fname == files[0]:
            # Insert page break at start of document
            pb_p = OxmlElement('w:p')
            pb_pPr = OxmlElement('w:pPr')
            pgBr = OxmlElement('w:pageBreakBefore')
            pgBr.set(qn('w:val'), '1')
            pb_pPr.append(pgBr)
            pb_p.append(pb_pPr)
            # Insert before first body element
            first_elem = src.element.body[0]
            src.element.body.insert(0, pb_p)

        composer.append(src)

composer.save(OUT_PATH)
size_kb = os.path.getsize(OUT_PATH) // 1024
print(f'\nSaved: {OUT_PATH}')
print(f'Size: {size_kb} KB')

# Verify images
import zipfile
with zipfile.ZipFile(OUT_PATH) as z:
    media = [f for f in z.namelist() if 'media' in f]
    print(f'Media files embedded: {len(media)}')
